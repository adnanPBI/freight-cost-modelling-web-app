from openpyxl import load_workbook
import json
rate='/mnt/data/Rate card example - FTL, LTL-PALLET and PALLET.xlsx'
alloc='/mnt/data/Allocation Key example.xlsx'
out={}
wb=load_workbook(rate,data_only=False,read_only=True)
out['rate_workbook']={s:{'rows':wb[s].max_row,'columns':wb[s].max_column} for s in wb.sheetnames}
wa=load_workbook(alloc,data_only=False,read_only=True)
ws=wa[wa.sheetnames[0]]
rows=[]
for r in ws.iter_rows(min_row=2,values_only=True):
    if r[1] and r[9]: rows.append({'dc':r[1],'mode':r[2],'dest_country':r[4],'dest_region':str(r[5]),'destination_zip':str(r[6]) if r[6] is not None else None,'customer_flag':r[7],'allocation_order':r[8],'carrier':r[9]})
out['allocation_rows']=rows
with open('/mnt/data/carrier-rate-allocation-platform/sample_import_evidence.json','w') as f: json.dump(out,f,indent=2)
print(json.dumps({'rate_sheets':out['rate_workbook'],'allocation_rows':len(rows)},indent=2))
